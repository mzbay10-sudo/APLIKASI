"""Uji alur aplikasi (launcher.py) dengan bot tiruan, tanpa Chrome/Flow.

Butuh tkinter + layar (Windows biasa sudah ada). Di Linux tanpa layar: xvfb-run.
"""

from __future__ import annotations

import json
import os
import runpy
import shutil
import sys
import tempfile
import unittest
from pathlib import Path
from unittest import mock

APP = Path(__file__).resolve().parents[1]

try:
    import tkinter  # noqa: F401
    HAS_TK = bool(os.environ.get("DISPLAY")) or sys.platform.startswith("win")
except ImportError:
    HAS_TK = False

# Bot tiruan: tiap run mengambil langkah berikutnya dari plan.json, menulis STATUS
# ke Excel, mencetak baris PROGRES, lalu keluar dengan kode yang ditentukan.
FAKE_WORKER = r'''
import json, sys
from pathlib import Path
from openpyxl import load_workbook
here = Path(__file__).parent
plan = json.loads((here / "plan.json").read_text())
count_file = here / "count.txt"
n = int(count_file.read_text()) if count_file.exists() else 0
count_file.write_text(str(n + 1))
step = plan[n]
args = sys.argv
profile = args[args.index("--profile-name") + 1]
with (here / "calls.txt").open("a") as f:
    f.write(profile + "\n")
excel = Path(args[args.index("--inputs") + 1])
wb = load_workbook(excel)
ws = wb.active
headers = [c.value for c in ws[1]]
if "STATUS" not in headers:
    ws.cell(row=1, column=len(headers) + 1, value="STATUS")
    headers.append("STATUS")
col = headers.index("STATUS") + 1
for i, (row, status) in enumerate(step.get("marks", {}).items(), start=1):
    print(f"2026-01-01 | INFO | PROGRES {i}/{len(step['marks'])} | {excel.name} | baris Excel {row}", flush=True)
    ws.cell(row=int(row), column=col, value=status)
wb.save(excel)
sys.exit(step["code"])
'''


@unittest.skipUnless(HAS_TK, "tkinter/layar tidak tersedia")
class LauncherFlowTest(unittest.TestCase):
    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp())
        for name in ("launcher.py", "bot.py", "capcut_plan.py"):
            shutil.copy2(APP / name, self.tmp / name)
        config = json.loads((APP / "config.example.json").read_text(encoding="utf-8"))
        (self.tmp / "config.json").write_text(json.dumps(config), encoding="utf-8")
        (self.tmp / "profiles.json").write_text(json.dumps({"A": "runtime/a", "B": "runtime/b"}), encoding="utf-8")
        (self.tmp / "profile_urls.json").write_text(json.dumps({
            "A": "https://labs.google/fx/tools/flow/project/aaaaaaaa-1111",
            "B": "https://labs.google/fx/tools/flow/project/bbbbbbbb-2222",
        }), encoding="utf-8")
        (self.tmp / "fake_worker.py").write_text(FAKE_WORKER, encoding="utf-8")
        (self.tmp / "data").mkdir()
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        ws.append(["Scene #", "Source Ref", "Prompt", "Character 1"])
        for i in range(1, 4):
            ws.append([i, f"¶{i}", f"Prompt {i}", "Xu Qing"])
        self.excel = self.tmp / "data" / "Episode_041.xlsx"
        wb.save(self.excel)

    def load_launcher(self):
        """Muat launcher.py salinan (APP_DIR = folder sementara, bot.py salinan)."""
        for name in ("bot", "capcut_plan"):
            sys.modules.pop(name, None)
        sys.path.insert(0, str(self.tmp))
        self.addCleanup(lambda: sys.path.remove(str(self.tmp)))
        self.addCleanup(lambda: [sys.modules.pop(n, None) for n in ("bot", "capcut_plan")])
        return runpy.run_path(str(self.tmp / "launcher.py"), run_name="launcher_test")

    def run_app(self, plan, rotate=False):
        (self.tmp / "plan.json").write_text(json.dumps(plan), encoding="utf-8")
        ns = self.load_launcher()
        fake = str(self.tmp / "fake_worker.py")
        ns["worker_command"].__globals__["worker_command"] = lambda extra: [sys.executable, fake, *extra]
        shown = []
        answers = mock.patch.multiple(
            "tkinter.messagebox",
            showinfo=lambda *a, **k: shown.append(("info", a)),
            showwarning=lambda *a, **k: shown.append(("warning", a)),
            showerror=lambda *a, **k: shown.append(("error", a)),
            askyesno=lambda *a, **k: True,
        )
        with answers:
            app = ns["FlowBotApp"]()
            app.profile_var.set("A")
            app.rotate_var.set(rotate)
            app.rotation_list.selection_set(0, "end")
            finished = []
            original = app._finished

            def finished_hook(code):
                original(code)
                if app.process is None and app.pending_launch is None and not app.job:
                    finished.append(code)
                    app.after(50, app.quit)

            app._finished = finished_hook
            app.after(100, app.start)
            app.after(60000, app.quit)  # pengaman
            app.mainloop()
            log = app.log.get("1.0", "end")
            status = app.status_var.get()
            app.destroy()
        calls = (self.tmp / "calls.txt").read_text().split() if (self.tmp / "calls.txt").exists() else []
        return calls, shown, log, status

    def test_auto_retry_failed_once_then_summary(self):
        plan = [
            {"code": 0, "marks": {"2": "SELESAI (2/2)", "3": "GAGAL (0/2) - error", "4": "SELESAI (2/2)"}},
            {"code": 0, "marks": {"3": "SELESAI (2/2)"}},
        ]
        calls, shown, log, _ = self.run_app(plan)
        self.assertEqual(calls, ["A", "A"])
        self.assertIn("ULANG OTOMATIS", log)
        self.assertIn("SELESAI 3", log)
        self.assertEqual(shown[-1][0], "info")

    def test_retry_only_once(self):
        plan = [
            {"code": 0, "marks": {"2": "SELESAI (2/2)", "3": "GAGAL (0/2)", "4": "SELESAI (2/2)"}},
            {"code": 0, "marks": {"3": "GAGAL (0/2)"}},
        ]
        calls, shown, log, _ = self.run_app(plan)
        self.assertEqual(calls, ["A", "A"])
        self.assertIn("GAGAL 1", log)

    def test_no_retry_when_everything_failed(self):
        plan = [{"code": 0, "marks": {"2": "GAGAL (0/2)", "3": "GAGAL (0/2)", "4": "GAGAL (0/2)"}}]
        calls, _, log, _ = self.run_app(plan)
        self.assertEqual(calls, ["A"])
        self.assertNotIn("ULANG OTOMATIS", log)

    def test_rotation_on_credit_limit(self):
        plan = [
            {"code": 9, "marks": {"2": "SELESAI (2/2)"}},
            {"code": 0, "marks": {"3": "SELESAI (2/2)", "4": "SELESAI (2/2)"}},
        ]
        calls, shown, log, _ = self.run_app(plan, rotate=True)
        self.assertEqual(calls, ["A", "B"])
        self.assertIn("PINDAH PROFIL OTOMATIS", log)
        rotation = json.loads((self.tmp / "profile_rotation.json").read_text())
        self.assertEqual(rotation["last_profile"], "B")

    def test_all_profiles_out_of_credit(self):
        plan = [{"code": 9, "marks": {}}, {"code": 10, "marks": {}}]
        calls, shown, _, _ = self.run_app(plan, rotate=True)
        self.assertEqual(calls, ["A", "B"])
        self.assertEqual(shown[-1][0], "warning")

    def test_eta_shown_in_status(self):
        ns = self.load_launcher()
        app = ns["FlowBotApp"]()
        try:
            with mock.patch.object(ns["time"], "monotonic", side_effect=[0.0, 60.0, 120.0, 180.0]):
                for n in (1, 2, 3, 4):
                    app.append_log(f"x | INFO | PROGRES {n}/10 | f.xlsx | prompt {n}/10 | baris Excel {n + 1}\n")
            self.assertIn("sisa ±7 menit", app.status_var.get())
        finally:
            app.destroy()


if __name__ == "__main__":
    unittest.main()
