"""Uji offline FlowBot (tanpa Google Flow / Chrome).

Jalankan dari folder aplikasi:
    .\\.venv\\Scripts\\python.exe -m unittest discover -s tests -v

Yang diuji: pembacaan Excel (format lama & STORY MODE), dry-run, penanda Excel
(termasuk saat Excel sedang dibuka), dan sinkron karakter antar profil dengan
profil Flow tiruan.
"""

from __future__ import annotations

import copy
import json
import os
import shutil
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path
from unittest import mock

APP = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(APP))

from openpyxl import Workbook, load_workbook  # noqa: E402

import bot  # noqa: E402

BASE_CONFIG = json.loads((APP / "config.example.json").read_text(encoding="utf-8"))


def make_config(tmp: Path) -> dict:
    config = copy.deepcopy(BASE_CONFIG)
    config["download"]["root_folder"] = str(tmp / "downloads")
    config["download"]["folder"] = str(tmp / "downloads" / "x")
    config["character_folder"] = str(tmp / "downloads" / "_KARAKTER")
    return config


def copy_bot(tmp: Path) -> Path:
    """Jalankan bot.py salinan agar log uji tidak masuk ke runtime/logs aplikasi."""
    shutil.copy2(APP / "bot.py", tmp / "bot.py")
    return tmp / "bot.py"


def write_positional(path: Path) -> None:
    """Format lama RG_xxx: kolom 1 prompt, kolom 2.. karakter."""
    wb = Workbook()
    ws = wb.active
    ws.append(["1", "2", "3", "4"])
    ws.append(["Prompt satu", "Xu Qing", "Wang Lin", 0])
    ws.append(["Prompt dua", "  wang   lin ", None, "None"])
    ws.append([None, "Hantu", None, None])  # prompt kosong -> dilewati
    ws.append(["Prompt tiga", "Hantu", "-", "N/A"])
    wb.save(path)


def write_story(path: Path) -> None:
    """Format STORY MODE: header bernama, prompt bukan di kolom pertama."""
    wb = Workbook()
    ws = wb.active
    ws.append(["Scene #", "Source Ref", "Prompt", "Character 1", "Character 2", "Key Artifact"])
    ws.append([1, "¶0", "Adegan satu", "Xu Qing", None, "Pedang"])
    ws.append([2, "¶1", "Adegan dua", "Li Mei", "Xu Qing", None])
    ws.append([3, "¶2", "Adegan tiga", None, None, None])
    wb.save(path)


class ExcelReadingTest(unittest.TestCase):
    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp())
        self.config = make_config(self.tmp)

    def test_positional_rows_and_characters(self):
        path = self.tmp / "RG_1-2.xlsx"
        write_positional(path)
        rows = bot.prompt_rows(path, self.config)
        self.assertEqual([n for n, _ in rows], [2, 3, 5])
        names = bot.excel_character_names([path], self.config)
        self.assertEqual(set(names), {"xu qing", "wang lin", "hantu"})

    def test_story_mode_uses_headers(self):
        path = self.tmp / "Episode_041.xlsx"
        write_story(path)
        rows = bot.prompt_rows(path, self.config)
        self.assertEqual(len(rows), 3)
        names = bot.excel_character_names([path], self.config)
        # Key Artifact tidak dihitung sebagai karakter.
        self.assertEqual(set(names), {"xu qing", "li mei"})
        file_config = bot.config_for_input(self.config, path)
        self.assertEqual(bot.scene_label(2, rows[0][1], file_config), "S001")

    def test_aliases_are_applied(self):
        path = self.tmp / "RG_1-2.xlsx"
        write_positional(path)
        self.config["character_aliases"] = {"Hantu": "Ghost King"}
        names = bot.excel_character_names([path], self.config)
        self.assertIn("ghost king", names)
        self.assertNotIn("hantu", names)

    def test_dry_run_cli(self):
        path = self.tmp / "Episode_041.xlsx"
        write_story(path)
        config_path = self.tmp / "config.json"
        config_path.write_text(json.dumps(self.config), encoding="utf-8")
        result = subprocess.run(
            [sys.executable, str(copy_bot(self.tmp)), "--config", str(config_path), "--dry-run", "--inputs", str(path)],
            capture_output=True, text=True, encoding="utf-8", timeout=120,
        )
        self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
        self.assertIn("3 baris berisi prompt", result.stderr)


class ExcelMarkerTest(unittest.TestCase):
    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp())
        self.config = make_config(self.tmp)
        self.path = self.tmp / "Episode_041.xlsx"
        write_story(self.path)
        self.file_config = bot.config_for_input(self.config, self.path)

    def test_marks_skip_finished_rows(self):
        marker = bot.ExcelMarker(self.path, self.file_config)
        marker.mark(2, ["Episode_041 S001-1.jpeg", "Episode_041 S001-2.jpeg"])
        marker.mark(3, [], "error")
        rows = bot.prompt_rows(self.path, self.config)
        self.assertEqual([n for n, _ in rows], [3, 4])  # SELESAI dilewati, GAGAL diulang
        backup = Path(self.file_config["download"]["folder"]) / "Episode_041_ASLI.xlsx"
        self.assertTrue(backup.is_file())

    def test_resume_when_excel_was_open(self):
        """Excel terbuka -> tanda masuk ke _TANDA.xlsx. Run berikutnya harus tetap
        melewati baris yang sudah SELESAI dan membawa tandanya ke file asli."""
        marker = bot.ExcelMarker(self.path, self.file_config)
        with mock.patch.object(bot.os, "replace", side_effect=PermissionError("locked")):
            marker.mark(2, ["Episode_041 S001-1.jpeg", "Episode_041 S001-2.jpeg"])
        fallback = Path(self.file_config["download"]["folder"]) / "Episode_041_TANDA.xlsx"
        self.assertTrue(fallback.is_file())
        rows = bot.prompt_rows(self.path, self.config)
        self.assertEqual([n for n, _ in rows], [3, 4])
        # Excel ditutup -> open_marker langsung menggabung _TANDA ke file asli.
        self.assertIsNotNone(bot.open_marker(self.path, self.file_config))
        self.assertFalse(fallback.exists())
        self.assertEqual([n for n, _ in bot.prompt_rows(self.path, self.config)], [3, 4])
        # Marker baru (Excel sudah ditutup) memindahkan tanda lama ke file asli.
        marker2 = bot.ExcelMarker(self.path, self.file_config)
        marker2.mark(3, ["Episode_041 S002-1.jpeg", "Episode_041 S002-2.jpeg"])
        sheet = load_workbook(self.path).active
        statuses = [sheet.cell(row=r, column=marker2.columns[bot.MARK_STATUS]).value for r in (2, 3, 4)]
        self.assertTrue(str(statuses[0]).startswith("SELESAI"))
        self.assertTrue(str(statuses[1]).startswith("SELESAI"))
        self.assertIsNone(statuses[2])
        self.assertFalse(fallback.exists())


# ---------------------------------------------------------------------------
# Sinkron karakter dengan profil Flow tiruan
# ---------------------------------------------------------------------------

class FakePage:
    def __init__(self, profile: str):
        self.profile = profile
        self.url = ""

    def set_default_timeout(self, _ms):
        pass

    def evaluate(self, *_args, **_kwargs):
        return True  # grid karakter termuat


class FakeContext:
    def __init__(self, profile: str):
        self.pages = [FakePage(profile)]
        self.closed = False

    def close(self):
        self.closed = True


class FakeFlow:
    """Isi karakter tiap profil + pencatat profil mana saja yang dibuka."""

    def __init__(self, tmp: Path, profiles: dict[str, list[str]], locked: set[str] = frozenset()):
        self.tmp = tmp
        self.chars = {name: list(chars) for name, chars in profiles.items()}
        self.locked = set(locked)
        self.opened: list[str] = []
        self.imported: list[tuple[str, str]] = []
        self.infos = []
        for name in profiles:
            folder = tmp / "profiles" / name
            folder.mkdir(parents=True, exist_ok=True)
            self.infos.append({"name": name, "dir": str(folder), "url": f"https://labs.google/fx/tools/flow/project/{name}-0000"})

    def profile_of(self, config):
        return str(config["url"]).rsplit("/", 1)[-1].split("-")[0]

    def open_context(self, _playwright, config):
        name = self.profile_of(config)
        self.opened.append(name)
        if name in self.locked:
            raise RuntimeError("Browser profile is already in use")
        return FakeContext(name)

    def open_characters_view(self, page, config):
        page.profile = self.profile_of(config)
        return "grid"

    def list_characters(self, page, _config=None, _mode="grid"):
        return [{"name": n, "src": f"https://img/{n}"} for n in self.chars[page.profile]]

    def export_characters(self, page, config, profile_name, only=None):
        folder = bot.character_folder(profile_name, config)
        folder.mkdir(parents=True, exist_ok=True)
        manifest = []
        for name in self.chars[page.profile]:
            if only is None or name.casefold() in only:
                file_name = bot.safe_file_name(name) + ".png"
                (folder / file_name).write_bytes(b"\x89PNG" + b"0" * 600)
                manifest.append({"name": name, "file": file_name, "personality": ""})
        (folder / "karakter.json").write_text(json.dumps({"profile": profile_name, "characters": manifest}), encoding="utf-8")
        return 0

    def import_characters(self, page, config, folder, names, report=None, deadline=None):
        for name in names:
            self.chars[page.profile].append(name)
            self.imported.append((folder.name, name))
            if report is not None:
                report.add(bot.norm_name(name))
        return 0

    def character_exists(self, page, _config, name):
        return any(bot.norm_name(c) == bot.norm_name(name) for c in self.chars[page.profile])

    def patches(self):
        return [
            mock.patch.object(bot, "load_profiles_info", lambda: self.infos),
            mock.patch.object(bot, "open_context", self.open_context),
            mock.patch.object(bot, "open_characters_view", self.open_characters_view),
            mock.patch.object(bot, "list_characters", self.list_characters),
            mock.patch.object(bot, "export_characters", self.export_characters),
            mock.patch.object(bot, "import_characters", self.import_characters),
            mock.patch.object(bot, "character_exists", self.character_exists),
        ]


class CharacterSyncTest(unittest.TestCase):
    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp())
        self.config = make_config(self.tmp)
        self.excel = self.tmp / "RG_1-2.xlsx"
        write_positional(self.excel)  # butuh: Xu Qing, Wang Lin, Hantu

    def run_sync(self, flow: FakeFlow, current: str = "A"):
        config = copy.deepcopy(self.config)
        config["url"] = next(i["url"] for i in flow.infos if i["name"] == current)
        config["profile_dir"] = next(i["dir"] for i in flow.infos if i["name"] == current)
        page = FakePage(current)
        with contextlib_exit(flow.patches()):
            bot.sync_characters_before_generate(object(), page, config, [self.excel], current)
        return config

    def test_missing_character_is_fetched_from_other_profile(self):
        flow = FakeFlow(self.tmp, {"A": ["Xu Qing"], "B": ["Wang Lin"], "C": ["Xu Qing"]})
        config = self.run_sync(flow)
        self.assertIn("Wang Lin", flow.chars["A"])
        self.assertEqual(flow.imported, [("B", "Wang Lin")])
        # Hantu tidak ada di profil mana pun -> dilewati saat generate.
        self.assertEqual(config["_missing_characters"], ["hantu"])
        # Katalog profil A ikut diperbarui dengan karakter yang baru dipindah.
        catalog = bot.load_catalog(config)
        self.assertIn("Wang Lin", catalog["A"]["names"])

    def test_locked_profile_is_skipped(self):
        flow = FakeFlow(self.tmp, {"A": [], "B": ["Wang Lin"], "C": ["Xu Qing", "Hantu"]}, locked={"B"})
        config = self.run_sync(flow)
        self.assertEqual(sorted(n for _, n in flow.imported), ["Hantu", "Xu Qing"])
        self.assertEqual(config["_missing_characters"], ["wang lin"])

    def test_saved_folder_used_before_opening_profiles(self):
        flow = FakeFlow(self.tmp, {"A": ["Xu Qing"], "B": ["Wang Lin", "Hantu"]})
        self.run_sync(flow)                     # run 1: buka B, unduh ke _KARAKTER/B
        flow2 = FakeFlow(self.tmp, {"A": ["Xu Qing"], "B": ["Wang Lin", "Hantu"]})
        self.run_sync(flow2)                    # run 2 (profil A kosong lagi): cukup dari folder
        self.assertEqual(flow2.opened, [])
        self.assertEqual(sorted(n for _, n in flow2.imported), ["Hantu", "Wang Lin"])

    def test_fresh_catalog_skips_profiles_without_character(self):
        flow = FakeFlow(self.tmp, {"A": ["Xu Qing", "Wang Lin"], "B": ["Lain"], "C": ["Lain"]})
        self.run_sync(flow)
        self.assertEqual(flow.opened, ["B", "C"])
        flow2 = FakeFlow(self.tmp, {"A": ["Xu Qing", "Wang Lin"], "B": ["Lain"], "C": ["Lain"]})
        self.run_sync(flow2)
        self.assertEqual(flow2.opened, [])     # katalog < 6 jam: B & C tidak punya Hantu

    def test_nothing_missing_opens_nothing(self):
        flow = FakeFlow(self.tmp, {"A": ["Xu Qing", "Wang Lin", "Hantu"], "B": ["Wang Lin"]})
        config = self.run_sync(flow)
        self.assertEqual(flow.opened, [])
        self.assertNotIn("_missing_characters", config)


@unittest.skipUnless(os.environ.get("FLOWBOT_SAMPLE_XLSX"), "set FLOWBOT_SAMPLE_XLSX=<Excel asli> untuk menguji file sungguhan")
class RealExcelTest(unittest.TestCase):
    """Uji Excel sungguhan (salinan; file asli tidak diubah)."""

    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp())
        self.config = make_config(self.tmp)
        source = Path(os.environ["FLOWBOT_SAMPLE_XLSX"])
        self.path = self.tmp / source.name
        shutil.copy2(source, self.path)

    def test_dry_run_and_marks(self):
        rows = bot.prompt_rows(self.path, self.config)
        self.assertGreater(len(rows), 0)
        config_path = self.tmp / "config.json"
        config_path.write_text(json.dumps(self.config), encoding="utf-8")
        result = subprocess.run(
            [sys.executable, str(copy_bot(self.tmp)), "--config", str(config_path), "--dry-run", "--inputs", str(self.path)],
            capture_output=True, text=True, encoding="utf-8", timeout=300,
        )
        self.assertEqual(result.returncode, 0, result.stderr[-2000:])
        file_config = bot.config_for_input(self.config, self.path)
        marker = bot.ExcelMarker(self.path, file_config)
        first, second = rows[0][0], rows[1][0]
        marker.mark(first, ["a-1.jpeg", "a-2.jpeg"])
        with mock.patch.object(bot.os, "replace", side_effect=PermissionError("locked")):
            marker.mark(second, ["b-1.jpeg"])
        left = [n for n, _ in bot.prompt_rows(self.path, self.config)]
        self.assertNotIn(first, left)
        self.assertNotIn(second, left)
        self.assertEqual(len(left), len(rows) - 2)
        # Sheet lain (README) tetap utuh.
        self.assertEqual(load_workbook(self.path).sheetnames, load_workbook(os.environ["FLOWBOT_SAMPLE_XLSX"]).sheetnames)


class contextlib_exit:
    """Masukkan beberapa mock.patch sekaligus."""

    def __init__(self, patches):
        self.patches = patches

    def __enter__(self):
        for patch in self.patches:
            patch.start()

    def __exit__(self, *exc):
        for patch in reversed(self.patches):
            patch.stop()
        return False


if __name__ == "__main__":
    unittest.main()
